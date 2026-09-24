from __future__ import annotations

import hashlib
import html
import io
import json
import re
from datetime import UTC, datetime
from typing import Any
from uuid import UUID

from docx import Document
from docx.shared import Pt
from ebooklib import epub  # type: ignore[import-untyped]
from pydantic import BaseModel, Field, model_validator
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from novel_writer.db.models import (
    ChapterRecord,
    ChapterRevisionRecord,
    ProjectRecord,
    QualityLabRecord,
    StateVersionRecord,
)
from novel_writer.domain.models import StoryState
from novel_writer.generation.archive import export_history, restore_history
from novel_writer.services.canonical_json import canonical_json_sha256
from novel_writer.services.creative_policies import CreativePolicy, CreativePolicyService
from novel_writer.services.errors import NotFoundError, WorkflowError
from novel_writer.services.style_profiles import StyleProfileService

ARCHIVE_FORMAT = "novel-writer-project-v1"
_CHAPTER_HEADING = re.compile(
    r"^#{1,6}\s*(?:第\s*([0-9一二三四五六七八九十百千零〇两]+)\s*章(?:\s*[：:]?\s*(.*))?|(.+))\s*$"
)


class PortableChapter(BaseModel):
    ordinal: int = Field(ge=1)
    title: str = Field(min_length=1, max_length=200)
    body: str = Field(min_length=1)


class ProjectArchive(BaseModel):
    format: str = ARCHIVE_FORMAT
    created_at: str
    project: dict[str, Any]
    formal_version: dict[str, Any]
    chapters: list[PortableChapter]
    style_profile: dict[str, Any] | None = None
    creative_policy: dict[str, Any] | None = None
    quality_regressions: list[dict[str, Any]] = Field(default_factory=list)
    generation_history: list[dict[str, Any]] = Field(default_factory=list)
    checksum_sha256: str = ""

    @model_validator(mode="after")
    def validate_archive(self) -> ProjectArchive:
        if self.format != ARCHIVE_FORMAT:
            raise ValueError("不支持的备份格式")
        ordinals = [item.ordinal for item in self.chapters]
        if len(ordinals) != len(set(ordinals)):
            raise ValueError("备份中存在重复章节序号")
        StoryState.model_validate(self.formal_version.get("state", {}))
        return self


def parse_manuscript(text: str) -> list[PortableChapter]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        raise WorkflowError("导入文件没有正文")
    matches: list[tuple[int, str, int]] = []
    for index, line in enumerate(normalized.splitlines()):
        match = _CHAPTER_HEADING.match(line.strip())
        if not match:
            continue
        number, explicit_title, generic_title = match.groups()
        if number:
            ordinal = _chinese_or_arabic_number(number)
            title = (explicit_title or f"第{ordinal}章").strip()
        else:
            ordinal = len(matches) + 1
            title = (generic_title or f"第{ordinal}章").strip()
        matches.append((ordinal, title, index))
    lines = normalized.splitlines()
    if not matches:
        return [PortableChapter(ordinal=1, title="第1章", body=normalized)]
    chapters: list[PortableChapter] = []
    for position, (ordinal, title, start) in enumerate(matches):
        end = matches[position + 1][2] if position + 1 < len(matches) else len(lines)
        body = "\n".join(lines[start + 1 : end]).strip()
        if body:
            chapters.append(PortableChapter(ordinal=ordinal, title=title, body=body))
    if not chapters:
        raise WorkflowError("识别到章节标题，但没有可导入的正文")
    return chapters


def archive_checksum(payload: dict[str, Any]) -> str:
    unsigned = dict(payload)
    unsigned["checksum_sha256"] = ""
    encoded = json.dumps(
        unsigned, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


class DataPortabilityService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session

    async def import_manuscript(self, title: str, text: str) -> dict[str, Any]:
        chapters = parse_manuscript(text)
        return await self._create_snapshot(title, StoryState(), chapters, "import")

    async def export_markdown(
        self,
        project_id: UUID,
        *,
        version_id: UUID | None = None,
        project_title: str | None = None,
    ) -> tuple[str, str]:
        project, version, chapters = await self._formal_snapshot(project_id, version_id)
        title = project_title or project.title
        parts = [f"# {title}"]
        for chapter in chapters:
            heading = f"第{chapter.ordinal}章"
            if chapter.title != heading:
                heading += f" {chapter.title}"
            parts.append(f"## {heading}\n\n{chapter.body}")
        return _safe_filename(title) + ".md", "\n\n".join(parts) + "\n"

    async def export_text(
        self,
        project_id: UUID,
        *,
        version_id: UUID | None = None,
        project_title: str | None = None,
    ) -> tuple[str, str]:
        project, _, chapters = await self._formal_snapshot(project_id, version_id)
        title = project_title or project.title
        parts = [title]
        for chapter in chapters:
            parts.append(f"第{chapter.ordinal}章 {chapter.title}\n\n{chapter.body}")
        return _safe_filename(title) + ".txt", "\n\n\n".join(parts) + "\n"

    async def export_docx(
        self,
        project_id: UUID,
        *,
        version_id: UUID | None = None,
        project_title: str | None = None,
    ) -> tuple[str, bytes]:
        project, _, chapters = await self._formal_snapshot(project_id, version_id)
        title = project_title or project.title
        document = Document()
        document.core_properties.title = title
        document.styles["Normal"].font.name = "宋体"
        document.styles["Normal"].font.size = Pt(12)
        document.sections[0].header.paragraphs[0].text = title
        for chapter in chapters:
            document.add_heading(f"第{chapter.ordinal}章 {chapter.title}", level=1)
            for paragraph in chapter.body.split("\n"):
                if paragraph.strip():
                    document.add_paragraph(paragraph)
        output = io.BytesIO()
        document.save(output)
        return _safe_filename(title) + ".docx", output.getvalue()

    async def export_epub(
        self,
        project_id: UUID,
        *,
        version_id: UUID | None = None,
        project_title: str | None = None,
    ) -> tuple[str, bytes]:
        project, _, chapters = await self._formal_snapshot(project_id, version_id)
        title = project_title or project.title
        book = epub.EpubBook()
        book.set_identifier(f"urn:novel-writer:{project.id}")
        book.set_title(title)
        book.set_language("zh-CN")
        chapter_items: list[Any] = []
        for index, chapter in enumerate(chapters, 1):
            item = epub.EpubHtml(
                title=f"第{chapter.ordinal}章 {chapter.title}",
                file_name=f"chapter-{index}.xhtml",
                lang="zh-CN",
            )
            body = "".join(
                f"<p>{html.escape(line)}</p>" for line in chapter.body.split("\n") if line.strip()
            )
            item.content = f"<h1>第{chapter.ordinal}章 {html.escape(chapter.title)}</h1>{body}"
            book.add_item(item)
            chapter_items.append(item)
        book.toc = tuple(chapter_items)
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        book.spine = ["nav", *chapter_items]
        output = io.BytesIO()
        epub.write_epub(output, book, {"epub3_pages": True})
        return _safe_filename(title) + ".epub", output.getvalue()

    async def create_archive(
        self,
        project_id: UUID,
        *,
        version_id: UUID | None = None,
        project_title: str | None = None,
    ) -> dict[str, Any]:
        project, version, chapters = await self._formal_snapshot(project_id, version_id)
        title = project_title or project.title
        style_profile = await StyleProfileService(self.session).get(project_id)
        creative_policy = await CreativePolicyService(self.session).get(project_id)
        regressions = (
            await self.session.scalars(
                select(QualityLabRecord)
                .where(
                    QualityLabRecord.project_id == project_id,
                    QualityLabRecord.kind == "regression",
                )
                .order_by(QualityLabRecord.created_at)
            )
        ).all()
        archive = ProjectArchive(
            generation_history=await export_history(self.session, project_id),
            created_at=datetime.now(UTC).isoformat(),
            project={"title": title, "source_project_id": str(project.id)},
            formal_version={"number": version.number, "state": version.state},
            chapters=chapters,
            style_profile={
                "selection_mode": style_profile.get("selection_mode", "unselected"),
                "genre_id": style_profile.get("genre_id", ""),
                "genre_card_id": style_profile.get("genre_card_id"),
                "secondary_genre_card_ids": style_profile.get("secondary_genre_card_ids", []),
                "assets": style_profile["assets"],
            },
            creative_policy=(
                {
                    key: value
                    for key, value in creative_policy.items()
                    if key not in {"project_id", "saved", "immutable_safety_boundaries"}
                }
                if creative_policy["saved"]
                else None
            ),
            quality_regressions=[
                {"name": item.name, "payload": item.payload} for item in regressions
            ],
        ).model_dump(mode="json")
        archive["checksum_sha256"] = archive_checksum(archive)
        return archive

    def inspect_archive(self, payload: dict[str, Any]) -> dict[str, Any]:
        archive = ProjectArchive.model_validate(payload)
        valid = archive.checksum_sha256 == archive_checksum(payload)
        return {
            "valid": valid,
            "format": archive.format,
            "title": archive.project.get("title", "未命名作品"),
            "formal_version": archive.formal_version.get("number"),
            "chapter_count": len(archive.chapters),
            "characters": sum(len(item.body) for item in archive.chapters),
            "quality_regression_count": len(archive.quality_regressions),
            "creative_policy_configured": archive.creative_policy is not None,
            "created_at": archive.created_at,
            "message": "备份完整，可以恢复为新项目。" if valid else "备份校验失败，请勿恢复。",
        }

    async def restore_archive(
        self, payload: dict[str, Any], title: str | None = None
    ) -> dict[str, Any]:
        inspection = self.inspect_archive(payload)
        if not inspection["valid"]:
            raise WorkflowError("备份校验失败，已阻止恢复")
        archive = ProjectArchive.model_validate(payload)
        restored_title = title or f"{archive.project.get('title', '未命名作品')}（恢复副本）"
        state = StoryState.model_validate(archive.formal_version["state"])
        result = await self._create_snapshot(restored_title, state, archive.chapters, "restore")
        if archive.style_profile:
            await StyleProfileService(self.session).save(
                UUID(result["project_id"]),
                list(archive.style_profile.get("assets", [])),
                str(archive.style_profile.get("genre_id", "")),
                selection_mode=(
                    "specified"
                    if archive.style_profile.get("selection_mode") == "specified"
                    else "unselected"
                ),
                genre_card_id=(
                    str(archive.style_profile["genre_card_id"])
                    if archive.style_profile.get("genre_card_id")
                    else None
                ),
                secondary_genre_card_ids=tuple(
                    str(card_id)
                    for card_id in archive.style_profile.get("secondary_genre_card_ids", [])
                ),
            )
        restored_project_id = UUID(result["project_id"])
        restored_project = await self.session.get(ProjectRecord, restored_project_id)
        assert restored_project is not None and restored_project.current_version_id is not None
        await restore_history(
            self.session,
            restored_project_id,
            restored_project.current_version_id,
            archive.generation_history,
        )
        if archive.creative_policy:
            await CreativePolicyService(self.session).save(
                restored_project_id, CreativePolicy.model_validate(archive.creative_policy)
            )
        for item in archive.quality_regressions:
            self.session.add(
                QualityLabRecord(
                    project_id=restored_project_id,
                    kind="regression",
                    name=str(item.get("name", "质量回归案例")),
                    status="ready",
                    payload=dict(item.get("payload", {})),
                    result=None,
                )
            )
        return result

    async def formal_snapshot_manifest(
        self,
        project_id: UUID,
        version_id: UUID | None = None,
    ) -> dict[str, Any]:
        project, version, _ = await self._formal_snapshot(project_id, version_id)
        chapter_records = list(
            await self.session.scalars(
                select(ChapterRecord).where(ChapterRecord.project_id == project.id)
            )
        )
        chapter_records.sort(key=lambda item: (item.display_ordinal or item.ordinal, item.ordinal))
        entries: list[dict[str, Any]] = []
        title_map = version.chapter_titles if isinstance(version.chapter_titles, dict) else {}
        for chapter in chapter_records:
            revision_id = version.chapter_revisions.get(str(chapter.id))
            if not revision_id:
                continue
            revision = await self.session.get(ChapterRevisionRecord, UUID(revision_id))
            if revision is None or revision.chapter_id != chapter.id:
                raise WorkflowError("formal snapshot revision is unavailable")
            entries.append(
                {
                    "chapter_id": str(chapter.id),
                    "revision_id": str(revision.id),
                    "ordinal": chapter.display_ordinal or chapter.ordinal,
                    "title": str(title_map.get(str(chapter.id)) or chapter.title),
                    "body_sha256": hashlib.sha256(revision.body.encode("utf-8")).hexdigest(),
                }
            )
        payload = {
            "schema_version": "formal-snapshot-manifest-v1",
            "project_id": str(project.id),
            "project_title": project.title,
            "version_id": str(version.id),
            "version_number": version.number,
            "state_sha256": canonical_json_sha256(version.state),
            "chapters": entries,
        }
        return {**payload, "snapshot_sha256": canonical_json_sha256(payload)}

    async def verify_formal_snapshot_manifest(
        self,
        expected: dict[str, Any],
    ) -> None:
        if expected.get("schema_version") != "formal-snapshot-manifest-v1":
            raise WorkflowError("formal snapshot manifest schema is invalid")
        try:
            project_id = UUID(str(expected["project_id"]))
            version_id = UUID(str(expected["version_id"]))
        except (KeyError, TypeError, ValueError) as error:
            raise WorkflowError("formal snapshot manifest identity is invalid") from error
        actual = await self.formal_snapshot_manifest(project_id, version_id)
        if actual != expected:
            raise WorkflowError("formal snapshot manifest changed")

    async def _formal_snapshot(
        self, project_id: UUID, version_id: UUID | None = None
    ) -> tuple[ProjectRecord, StateVersionRecord, list[PortableChapter]]:
        project = await self.session.get(ProjectRecord, project_id)
        if project is None or project.current_version_id is None:
            raise NotFoundError("project not found")
        version = await self.session.get(
            StateVersionRecord, version_id or project.current_version_id
        )
        if version is None or version.project_id != project.id:
            raise NotFoundError("formal version not found")
        chapter_records = list(
            await self.session.scalars(
                select(ChapterRecord).where(ChapterRecord.project_id == project.id)
            )
        )
        chapter_records.sort(key=lambda item: (item.display_ordinal or item.ordinal, item.ordinal))
        chapters: list[PortableChapter] = []
        title_map = version.chapter_titles if isinstance(version.chapter_titles, dict) else {}
        for chapter in chapter_records:
            revision_id = version.chapter_revisions.get(str(chapter.id))
            if not revision_id:
                continue
            revision = await self.session.get(ChapterRevisionRecord, UUID(revision_id))
            if revision is not None:
                chapters.append(
                    PortableChapter(
                        ordinal=chapter.display_ordinal or chapter.ordinal,
                        title=str(title_map.get(str(chapter.id)) or chapter.title),
                        body=revision.body,
                    )
                )
        return project, version, chapters

    async def _create_snapshot(
        self,
        title: str,
        state: StoryState,
        chapters: list[PortableChapter],
        source: str,
    ) -> dict[str, Any]:
        project = ProjectRecord(title=title)
        self.session.add(project)
        await self.session.flush()
        version = StateVersionRecord(
            project_id=project.id,
            number=1,
            state=state.model_dump(mode="json"),
            chapter_revisions={},
            chapter_titles={},
        )
        self.session.add(version)
        await self.session.flush()
        revision_map: dict[str, str] = {}
        title_map: dict[str, str] = {}
        chapter_ids: dict[int, str] = {}
        for item in sorted(chapters, key=lambda value: value.ordinal):
            chapter = ChapterRecord(
                project_id=project.id,
                ordinal=item.ordinal,
                display_ordinal=item.ordinal,
                title=item.title,
            )
            self.session.add(chapter)
            await self.session.flush()
            revision = ChapterRevisionRecord(
                chapter_id=chapter.id,
                base_version_id=version.id,
                body=item.body,
                status="committed",
            )
            self.session.add(revision)
            await self.session.flush()
            chapter.current_revision_id = revision.id
            revision_map[str(chapter.id)] = str(revision.id)
            title_map[str(chapter.id)] = chapter.title
            chapter_ids[item.ordinal] = str(chapter.id)
        version.state = StoryState.model_validate(
            _rebind_chapter_references(state.model_dump(mode="json"), chapter_ids)
        ).model_dump(mode="json")
        version.chapter_revisions = revision_map
        version.chapter_titles = title_map
        project.current_version_id = version.id
        return {
            "project_id": str(project.id),
            "version_id": str(version.id),
            "version": 1,
            "chapter_count": len(chapters),
            "source": source,
        }


def _rebind_chapter_references(value: Any, chapters: dict[int, str]) -> Any:
    """Copied formal text gets new IDs; historical generation receipts remain untouched."""
    if isinstance(value, list):
        return [_rebind_chapter_references(item, chapters) for item in value]
    if isinstance(value, dict):
        result = {key: _rebind_chapter_references(item, chapters) for key, item in value.items()}
        if "chapter_id" in value and value.get("chapter_ordinal") in chapters:
            result["chapter_id"] = chapters[value["chapter_ordinal"]]
        return result
    return value


def _safe_filename(title: str) -> str:
    cleaned = re.sub(r"[\\/:*?\"<>|]+", "_", title).strip(" .")
    return cleaned or "小说"


def _chinese_or_arabic_number(value: str) -> int:
    if value.isdigit():
        return int(value)
    digits = {
        "零": 0,
        "〇": 0,
        "一": 1,
        "二": 2,
        "两": 2,
        "三": 3,
        "四": 4,
        "五": 5,
        "六": 6,
        "七": 7,
        "八": 8,
        "九": 9,
    }
    units = {"十": 10, "百": 100, "千": 1000}
    total = current = 0
    for character in value:
        if character in digits:
            current = digits[character]
        elif character in units:
            total += (current or 1) * units[character]
            current = 0
    return total + current
