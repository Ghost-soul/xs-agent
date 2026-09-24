import io
import zipfile
from copy import deepcopy
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4
from xml.etree import ElementTree

import pytest
from docx import Document

from novel_writer.services.data_portability import (
    ARCHIVE_FORMAT,
    DataPortabilityService,
    PortableChapter,
    archive_checksum,
    parse_manuscript,
)


def test_parse_markdown_with_arabic_and_chinese_ordinals() -> None:
    chapters = parse_manuscript(
        "# 小说名\n\n## 第一章 入城\n\n正文一\n\n## 第2章 夜雨\n\n正文二"
    )

    assert [(item.ordinal, item.title, item.body) for item in chapters] == [
        (1, "入城", "正文一"),
        (2, "夜雨", "正文二"),
    ]


def test_plain_text_without_headings_imports_as_one_chapter() -> None:
    chapters = parse_manuscript("风从城门吹进来。\n守军举起了灯。")

    assert len(chapters) == 1
    assert chapters[0].title == "第1章"


def test_archive_checksum_detects_changed_body() -> None:
    payload = {
        "format": ARCHIVE_FORMAT,
        "created_at": "2026-07-22T00:00:00+00:00",
        "project": {"title": "测试"},
        "formal_version": {"number": 3, "state": {}},
        "chapters": [{"ordinal": 1, "title": "第一章", "body": "原文"}],
        "checksum_sha256": "",
    }
    payload["checksum_sha256"] = archive_checksum(payload)
    changed = deepcopy(payload)
    changed["chapters"][0]["body"] = "被修改"

    service = DataPortabilityService(None)  # type: ignore[arg-type]
    assert service.inspect_archive(payload)["valid"] is True
    assert service.inspect_archive(changed)["valid"] is False


def test_empty_manuscript_is_rejected() -> None:
    with pytest.raises(Exception, match="没有正文"):
        parse_manuscript("   ")


@pytest.mark.asyncio
async def test_docx_export_is_an_independently_parseable_openxml_package() -> None:
    project = SimpleNamespace(id=uuid4(), title="长夜 & 火")
    chapters = [
        PortableChapter(
            ordinal=1,
            title="门 <开>",
            body="第一段 & 回声\n第二段 <落幕>",
        )
    ]
    service = DataPortabilityService(None)  # type: ignore[arg-type]
    service._formal_snapshot = AsyncMock(  # type: ignore[method-assign]
        return_value=(project, SimpleNamespace(number=1), chapters)
    )

    filename, payload = await service.export_docx(project.id)

    assert filename == "长夜 & 火.docx"
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        assert archive.testzip() is None
        assert "word/document.xml" in archive.namelist()
        ElementTree.fromstring(archive.read("word/document.xml"))
    document = Document(io.BytesIO(payload))
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "第1章 门 <开>",
        "第一段 & 回声",
        "第二段 <落幕>",
    ]
    assert document.sections[0].header.paragraphs[0].text == "长夜 & 火"


@pytest.mark.asyncio
async def test_epub_export_has_valid_zip_xml_manifest_spine_and_navigation() -> None:
    project = SimpleNamespace(id=uuid4(), title="长夜 & 火")
    chapters = [
        PortableChapter(ordinal=1, title="门 <开>", body="第一段 & 回声"),
        PortableChapter(ordinal=2, title="归途", body="第二段 <落幕>"),
    ]
    service = DataPortabilityService(None)  # type: ignore[arg-type]
    service._formal_snapshot = AsyncMock(  # type: ignore[method-assign]
        return_value=(project, SimpleNamespace(number=1), chapters)
    )

    filename, payload = await service.export_epub(project.id)

    assert filename == "长夜 & 火.epub"
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        assert archive.testzip() is None
        first = archive.infolist()[0]
        assert first.filename == "mimetype"
        assert first.compress_type == zipfile.ZIP_STORED
        assert archive.read("mimetype") == b"application/epub+zip"
        expected_xml = {
            "META-INF/container.xml",
            "EPUB/content.opf",
            "EPUB/nav.xhtml",
            "EPUB/chapter-1.xhtml",
            "EPUB/chapter-2.xhtml",
        }
        assert expected_xml.issubset(archive.namelist())
        parsed = {
            path: ElementTree.fromstring(archive.read(path)) for path in expected_xml
        }

    container_ns = {"c": "urn:oasis:names:tc:opendocument:xmlns:container"}
    rootfile = parsed["META-INF/container.xml"].find(".//c:rootfile", container_ns)
    assert rootfile is not None
    assert rootfile.attrib["full-path"] == "EPUB/content.opf"

    opf_ns = {"opf": "http://www.idpf.org/2007/opf"}
    manifest_hrefs = {
        item.attrib["href"]
        for item in parsed["EPUB/content.opf"].findall(".//opf:manifest/opf:item", opf_ns)
    }
    spine_refs = [
        item.attrib["idref"]
        for item in parsed["EPUB/content.opf"].findall(".//opf:spine/opf:itemref", opf_ns)
    ]
    assert {"chapter-1.xhtml", "chapter-2.xhtml", "nav.xhtml"}.issubset(manifest_hrefs)
    assert spine_refs[-2:] == ["chapter_0", "chapter_1"]

    xhtml_ns = {"x": "http://www.w3.org/1999/xhtml"}
    links = parsed["EPUB/nav.xhtml"].findall(".//x:nav/x:ol/x:li/x:a", xhtml_ns)
    assert [(item.attrib["href"], item.text) for item in links] == [
        ("chapter-1.xhtml", "第1章 门 <开>"),
        ("chapter-2.xhtml", "第2章 归途"),
    ]
    first_chapter = parsed["EPUB/chapter-1.xhtml"]
    assert first_chapter.findtext(".//x:h1", namespaces=xhtml_ns) == "第1章 门 <开>"
    assert first_chapter.findtext(".//x:p", namespaces=xhtml_ns) == "第一段 & 回声"
